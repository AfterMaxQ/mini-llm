"""从实验主笔记生成 Word，并用本机 Word 渲染页面。"""
import argparse
import json
import re
import shutil
import time
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from lab import ROOT, now, sha256, write_json

VOLUMES = {"01": ("环境与训练机制", range(0, 8)), "02": ("微调与数据", range(8, 15)),
           "03": ("蒸馏", range(15, 19)), "04": ("推理与 Agent 评测", range(19, 26))}


def text(paragraph, content):
    for part in re.split(r"(`[^`]+`)", content):
        run = paragraph.add_run(part.strip("`") if part.startswith("`") else part)
        if part.startswith("`"):
            run.font.name = "Consolas"
            run.font.size = Pt(10)


def styles(document, title):
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.2)
    section.left_margin = section.right_margin = Cm(2)
    for name, size in [("Normal", 11), ("Title", 20), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)]:
        style = document.styles[name]
        style.font.name = "Calibri"
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string("1F2937" if name == "Normal" else "17365D")
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.25
        if name.startswith("Heading"):
            style.paragraph_format.space_before = Pt(8)
            style.paragraph_format.line_spacing = 1.1
    section.header.paragraphs[0].text = title
    code = document.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = "Consolas"
    code.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
    code.font.size = Pt(9)
    code.paragraph_format.line_spacing = 1.05
    code.paragraph_format.space_before = code.paragraph_format.space_after = Pt(6)
    code.paragraph_format.left_indent = code.paragraph_format.right_indent = Cm(0.2)
    code.paragraph_format.keep_together = True
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def add_markdown(document, path, page_break_before=False):
    first_paragraph = len(document.paragraphs)
    lines = path.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("```"):
            block = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index]); index += 1
            paragraph = document.add_paragraph(style="Code Block")
            paragraph.add_run("\n".join(block))
            shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "F3F4F6")
            paragraph._p.get_or_add_pPr().append(shade)
            index += 1
            continue
        if line.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                values = [x.strip().strip("`") for x in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", x) for x in values):
                    rows.append(values)
                index += 1
            table = document.add_table(rows=0, cols=len(rows[0]))
            table.style = "Table Grid"
            for number, values in enumerate(rows):
                row = table.add_row()
                row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                cells = row.cells
                for cell, value in zip(cells, values):
                    text(cell.paragraphs[0], value)
                    cell.paragraphs[0].paragraph_format.space_after = Pt(2)
                    cell.paragraphs[0].paragraph_format.line_spacing = 1.0
                if number == 0:
                    properties = table.rows[0]._tr.get_or_add_trPr()
                    properties.append(OxmlElement("w:tblHeader"))
                    for cell in cells:
                        shade = OxmlElement("w:shd")
                        shade.set(qn("w:fill"), "F3F4F6")
                        cell._tc.get_or_add_tcPr().append(shade)
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            continue
        picture = re.fullmatch(r"!\[([^]]*)\]\(([^)]+)\)", line)
        if picture:
            image = (path.parent / picture.group(2)).resolve()
            paragraph = document.add_paragraph()
            paragraph.alignment = 1
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.keep_with_next = True
            paragraph.add_run().add_picture(str(image), width=Cm(15))
            caption = document.add_paragraph(picture.group(1))
            caption.alignment = 1
            index += 1
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)", line)
        if heading:
            document.add_heading(heading.group(2), level=len(heading.group(1)))
        else:
            text(document.add_paragraph(), line)
        index += 1
    # 标题随正文一起换页，避免独立分页段落被挤出一张空白页。
    if page_break_before and len(document.paragraphs) > first_paragraph:
        document.paragraphs[first_paragraph].paragraph_format.page_break_before = True


def render(path, directory):
    import fitz
    import psutil
    import win32com.client
    import win32process
    existing = {p.pid for p in psutil.process_iter(["name"]) if (p.info["name"] or "").lower() == "winword.exe"}
    directory.mkdir(parents=True, exist_ok=True)
    pdf = directory / (path.stem + ".pdf")
    application = win32com.client.DispatchEx("Word.Application")
    application.Visible = False
    application.DisplayAlerts = 0
    document = None
    owned_pid = None
    try:
        document = application.Documents.Open(str(path.resolve()), ReadOnly=True, AddToRecentFiles=False)
        owned_pid = win32process.GetWindowThreadProcessId(document.ActiveWindow.Hwnd)[1]
        document.Repaginate()
        document.ExportAsFixedFormat(str(pdf.resolve()), 17, OpenAfterExport=False)
    finally:
        if document is not None:
            document.Close(False)
        if owned_pid is None:
            created = {p.pid for p in psutil.process_iter(["name"]) if (p.info["name"] or "").lower() == "winword.exe"} - existing
            if len(created) == 1:
                owned_pid = created.pop()
        if owned_pid is not None and owned_pid not in existing:
            application.Quit()
    pages = []
    with fitz.open(pdf) as document:
        for index, page in enumerate(document):
            image = directory / f"page-{index + 1:02d}.png"
            page.get_pixmap(dpi=120).save(image)
            pages.append({"page": index + 1, "png": str(image), "text_characters": len(page.get_text())})
    # 分册缩短后，移除这次 PDF 已不存在的旧页面图片。
    for old in directory.glob('page-*.png'):
        if int(old.stem.split('-')[-1]) > len(pages):
            old.unlink()
    write_json(directory / "render.json", {"word_sha256": sha256(path), "rendered": now(),
                                           "pdf_sha256": sha256(pdf), "visual_checked": False, "pages": pages})
    print(json.dumps({"word": str(path), "pdf": str(pdf), "pages": pages}, ensure_ascii=False), flush=True)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--volume", choices=VOLUMES, default="01")
    parser.add_argument("--render", action="store_true")
    args = parser.parse_args()
    title, numbers = VOLUMES[args.volume]
    notes = [ROOT / f"experiments/E{number:02d}/notes.md" for number in numbers]
    notes = [p for p in notes if p.exists()]
    if not notes:
        raise RuntimeError("该分册没有真实实验笔记")
    document = Document()
    styles(document, title)
    document.add_heading(f"MiniLLM 实验册 {args.volume}：{title}", 0)
    document.add_paragraph(f"更新：{now()}。本册按实际实验整理，原始运行记录单独保存。")
    for index, note in enumerate(notes):
        add_markdown(document, note, page_break_before=bool(index))
    summary = ROOT / f"docs/reports/summaries/{args.volume}.md"
    if summary.exists():
        add_markdown(document, summary)
    directory = ROOT / "docs/reports"
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / f"{args.volume}-{title}.docx"
    staging = ROOT / ".local/report-staging" / f"{time.time_ns()}-{path.name}"
    staging.parent.mkdir(parents=True, exist_ok=True)
    document.save(staging)
    snapshots = ROOT / ".local/doc-snapshots" / args.volume
    snapshots.mkdir(parents=True, exist_ok=True)
    stamp = time.time_ns()
    if path.exists():
        shutil.copy2(path, snapshots / f"{stamp}-previous.docx")
    try:
        staging.replace(path)
    except PermissionError:
        path = snapshots / f"{stamp}-locked.docx"
        staging.replace(path)
    # 只清理本分册的临时快照，不触及实验原始数据。
    for old in sorted(snapshots.glob("*.docx"), key=lambda p: p.stat().st_mtime, reverse=True)[5:]:
        old.unlink()
    print(f"Word 已保存：{path}", flush=True)
    if path.parent==directory:
        receipt=directory/f'visual-{args.volume}.json'
        previous=json.loads(receipt.read_text(encoding='utf-8')) if receipt.exists() else {}
        current_hash=sha256(path)
        if previous.get('word_sha256')!=current_hash:
            checked=previous if previous.get('visual_checked') else previous.get('previous_checked',{})
            write_json(receipt,{'word_sha256':current_hash,'visual_checked':False,
                               'updated':now(),'previous_checked':checked})
    from index import main as update_index
    update_index()
    if args.render:
        render(path, ROOT / ".local/render" / args.volume)


if __name__ == "__main__":
    main()
